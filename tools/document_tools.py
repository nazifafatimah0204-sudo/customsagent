import re
from pathlib import Path
from typing import List, Dict
from pydantic import Field
from crewai.tools import BaseTool
from pypdf import PdfReader
from docx import Document


def read_text_file(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix}")


class DocumentReaderTool(BaseTool):
    name: str = "Document Reader"
    description: str = "Read a PDF, DOCX, or TXT shipment document using its temporary file path. Input must be a file_path."
    def _run(self, file_path: str) -> str:
        try:
            content = read_text_file(file_path)
            if not content.strip():
                return "No selectable text found. The PDF may be scanned; OCR is required."
            return content[:24000]
        except Exception as exc:
            return f"Document read error: {exc}"


class CrossDocumentCompareTool(BaseTool):
    name: str = "Cross Document Compare"
    description: str = "Compare invoice and packing-list text for numeric, date, currency, HS-code, quantity, and weight discrepancies. Provide invoice_text and packing_list_text."
    def _run(self, invoice_text: str, packing_list_text: str) -> str:
        patterns = {
            "HS-code candidates": r"\b(?:HS(?:\s*Code)?|H\.S\.)\s*[:#-]?\s*([0-9][0-9 .-]{3,12})",
            "dates": r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b",
            "currencies / amounts": r"\b(?:USD|US\$|EUR|GBP|PKR|CNY|AED|SAR)\s?[\d,]+(?:\.\d{1,2})?\b",
            "weights": r"\b\d+(?:,\d{3})*(?:\.\d+)?\s?(?:kg|kgs|kilograms?|g|grams?|mt|tonnes?)\b",
            "quantities": r"\b\d+(?:,\d{3})*(?:\.\d+)?\s?(?:pcs?|pieces?|units?|cartons?|boxes?|sets?)\b",
        }
        out = ["Deterministic text scan (candidate matches only; verify context manually):"]
        for label, pattern in patterns.items():
            inv = re.findall(pattern, invoice_text, flags=re.I)
            pack = re.findall(pattern, packing_list_text, flags=re.I)
            out.append(f"\n{label}\n- Invoice: {inv[:30] or 'No candidate found'}\n- Packing list: {pack[:30] or 'No candidate found'}")
        out.append("\nThis tool reports candidates, not proof of a mismatch. Compare the surrounding labels and units.")
        return "\n".join(out)


class RegulationSearchTool(BaseTool):
    name: str = "Regulation Library Search (RAG)"
    description: str = "Search uploaded regulation and guidance text for relevant passages. Input a plain-language query."
    regulation_texts: List[Dict[str, str]] = Field(default_factory=list)

    def _run(self, query: str) -> str:
        terms = {t.lower() for t in re.findall(r"[a-zA-Z0-9]{3,}", query)}
        chunks = []
        for doc in self.regulation_texts:
            text = doc.get("text", "")
            paragraphs = [p.strip() for p in re.split(r"\n{1,}|(?<=[.!?])\s+", text) if p.strip()]
            for i, paragraph in enumerate(paragraphs):
                words = {t.lower() for t in re.findall(r"[a-zA-Z0-9]{3,}", paragraph)}
                score = len(terms & words)
                if score:
                    chunks.append((score, doc.get("name", "uploaded source"), paragraph))
        chunks.sort(key=lambda item: item[0], reverse=True)
        if not chunks:
            return "No matching passage found in the supplied regulation files. Do not infer a legal requirement from this absence."
        selected = chunks[:8]
        return "\n\n".join(
            f"[Source: {name} | keyword relevance: {score}]\n{paragraph[:1400]}"
            for score, name, paragraph in selected
        )


class ReportStructureTool(BaseTool):
    name: str = "Compliance Report Structure"
    description: str = "Organize draft findings into a clear report outline. Input the findings text."
    def _run(self, findings: str) -> str:
        return (
            "Use this report structure:\n"
            "1. Executive summary\n"
            "2. High-priority issues\n"
            "3. Medium-priority issues\n"
            "4. Low-priority / verification notes\n"
            "5. Cross-document evidence\n"
            "6. Regulation passages and source filenames\n"
            "7. Missing information and recommended human follow-up\n"
            "8. Limitations and explicit statement that this is not an official customs ruling\n\n"
            "Draft findings to organize:\n" + findings[:16000]
        )
